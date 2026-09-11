import io
from datetime import datetime

import arabic_reshaper
from bidi.algorithm import get_display
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from fpdf import FPDF

ARABIC_RANGES = (
    (0x0600, 0x06FF),
    (0x0750, 0x077F),
    (0x08A0, 0x08FF),
    (0xFB50, 0xFDFF),
    (0xFE70, 0xFEFF),
)

# Polices Windows couvrant à la fois le latin (français/anglais) et l'arabe.
PDF_FONT_CANDIDATES = [
    (r"C:\Windows\Fonts\tahoma.ttf", r"C:\Windows\Fonts\tahomabd.ttf"),
    (r"C:\Windows\Fonts\arialuni.ttf", r"C:\Windows\Fonts\arialuni.ttf"),
]


def contains_arabic(text: str) -> bool:
    return any(any(lo <= ord(ch) <= hi for lo, hi in ARABIC_RANGES) for ch in text)


def _shape_line(line: str) -> str:
    if not line.strip():
        return ""
    return get_display(arabic_reshaper.reshape(line))


def _set_paragraph_rtl(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))


def build_docx(messages: list[dict], sources: list[str]) -> bytes:
    doc = Document()
    doc.add_heading("Conversation — Assistant Articles", level=1)

    meta = doc.add_paragraph(f"Exporté le {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    meta.runs[0].italic = True

    if sources:
        doc.add_paragraph("Articles utilisés : " + ", ".join(sources))

    doc.add_paragraph()

    for msg in messages:
        label = "Vous" if msg["role"] == "user" else "Assistant"
        content = msg["content"]
        is_rtl = contains_arabic(content)

        p_label = doc.add_paragraph()
        run_label = p_label.add_run(f"{label} :")
        run_label.bold = True

        p_content = doc.add_paragraph(content)
        if is_rtl:
            for paragraph in (p_label, p_content):
                _set_paragraph_rtl(paragraph)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in paragraph.runs:
                    run.font.rtl = True

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _register_pdf_font(pdf: FPDF) -> str:
    for regular, bold in PDF_FONT_CANDIDATES:
        try:
            pdf.add_font("ChatFont", "", regular)
            pdf.add_font("ChatFont", "B", bold)
            return "ChatFont"
        except (RuntimeError, FileNotFoundError):
            continue
    return "helvetica"


def _write_line(pdf: FPDF, h: float, text: str, align: str = "L") -> None:
    # new_x/new_y ramènent le curseur en marge gauche : par défaut fpdf2 le laisse
    # à droite après un multi_cell, ce qui épuise la largeur disponible ligne après ligne.
    pdf.multi_cell(0, h, text, align=align, new_x="LMARGIN", new_y="NEXT")


def build_pdf(messages: list[dict], sources: list[str]) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    font = _register_pdf_font(pdf)

    pdf.set_font(font, "B", 16)
    _write_line(pdf, 10, "Conversation - Assistant Articles", align="C")
    pdf.ln(2)

    pdf.set_font(font, "", 9)
    _write_line(pdf, 6, f"Exporté le {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    if sources:
        _write_line(pdf, 6, "Articles utilisés : " + ", ".join(sources))
    pdf.ln(4)

    for msg in messages:
        label = "Vous" if msg["role"] == "user" else "Assistant"
        pdf.set_font(font, "B", 11)
        _write_line(pdf, 7, _shape_line(f"{label} :"))
        pdf.set_font(font, "", 11)
        for line in msg["content"].split("\n"):
            _write_line(pdf, 7, _shape_line(line))
        pdf.ln(3)

    return bytes(pdf.output())
